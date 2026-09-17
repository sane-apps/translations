"""DOCX helpers for Logos Personal Books."""
from __future__ import annotations

import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


def setup_document(
    *,
    title: str,
    author: str,
    subject: str = "",
    keywords: str = "",
    comments: str = "",
) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    sizes = {
        "Title": 28,
        "Subtitle": 15,
        "Heading 1": 19,
        "Heading 2": 15,
        "Heading 3": 12,
        "Caption": 10,
    }
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"]:
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.size = Pt(sizes.get(name, 11))
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.widow_control = True
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        for fonts in style.element.iter(qn("w:rFonts")):
            for attr in list(fonts.attrib):
                if attr.endswith("Theme"):
                    del fonts.attrib[attr]
    normal = doc.styles["Normal"]
    normal.paragraph_format.line_spacing = 1.08
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[name].paragraph_format.keep_with_next = True
        doc.styles[name].paragraph_format.space_before = Pt(12)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "en-US")
    normal.element.get_or_add_rPr().append(lang)
    props = doc.core_properties
    props.title = title
    props.author = author
    props.subject = subject
    props.keywords = keywords
    props.comments = comments
    return doc


class BookmarkStore:
    def __init__(self):
        self.ids: dict[str, str] = {}

    def add(self, paragraph, name: str):
        assert name not in self.ids, name
        number = str(len(self.ids) + 1)
        self.ids[name] = number
        begin = OxmlElement("w:bookmarkStart")
        begin.set(qn("w:id"), number)
        begin.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), number)
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, begin)
        paragraph._p.append(end)


def hyperlink(paragraph, text: str, target: str, *, internal: bool = False):
    h = OxmlElement("w:hyperlink")
    if internal:
        h.set(qn("w:anchor"), target)
    else:
        h.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    r = OxmlElement("w:r")
    pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "174A75")
    pr.append(c)
    r.append(pr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def logos_safe_headword(text: str) -> str:
    """Logos PBB rejects Headwords that contain colons (they are ignored at compile)."""
    hw = (text or "").replace("\n", " ").strip()
    hw = hw.replace(":", " —")
    hw = re.sub(r"\s+", " ", hw)
    # Keep ArticleCache usable; truncate pathological ANF argument titles.
    if len(hw) > 120:
        hw = hw[:117].rstrip(" —,;.") + "…"
    return hw


def add_heading_with_headword(
    doc: Document,
    bookmarks: BookmarkStore,
    text: str,
    level: int,
    key: str,
    *,
    headword_source: str | None = None,
):
    """Clean heading text + separate Headword paragraph (Logos ArticleCache-safe).

    ``text`` is the displayed heading — plain dot-form verse refs only, never
    ``[[… >> Bible:…]]`` (headings with Bible links warn at compile; see
    ``pipeline.bible_links.link_heading_verses``). The Headword milestone is
    built from ``headword_source`` (default: ``text``) so display text never
    leaks into ArticleCache titles.
    """
    p = doc.add_heading(text, level)
    bookmarks.add(p, key)
    doc.add_paragraph(f"[[@Headword:{logos_safe_headword(headword_source or text)}]]")
    return p


def assert_internal_links(doc: Document, bookmark_ids: dict):
    for h in doc.element.iter(qn("w:hyperlink")):
        anchor = h.get(qn("w:anchor"))
        assert anchor is None or anchor in bookmark_ids, anchor
